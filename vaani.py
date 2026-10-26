#!/usr/bin/env python3
"""
Vaani Lead Engine v3.12
======================
Type a location (even a small area like "Gota, Ahmedabad") and get 4 lead
databases about restaurants there:

    1. restaurant_names.csv   - every restaurant name found
    2. locations.csv          - their addresses
    3. phone_numbers.csv      - MOBILE numbers only (landlines are dropped)
    4. email_ids.csv          - email addresses found on their websites

Plus a bonus: all_leads.xlsx - everything combined in one Excel file.

The scraper is BUILT IN - this one file does everything. No Docker, no
separate engine, no API keys. It opens Google Maps in a real browser
window (don't close it while it runs), reads each restaurant, then checks
the restaurant's own website for an email address.

Needs: Python + the two helpers in requirements.txt
    py -m pip install -r requirements.txt
"""
import argparse
import csv
import os
import random
import re
import subprocess
import sys
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from datetime import date
import sqlite3

HERE = os.path.dirname(os.path.abspath(__file__))

VERSION = "3.12.0"
UPDATE_URL = ("https://raw.githubusercontent.com/yash6969-io/"
              "vaani-lead-engine/main/vaani.py")


# ---------------------------------------------------------------- auto-update

def _version_tuple(text):
    try:
        return tuple(int(x) for x in text.strip().split("."))
    except Exception:
        return (0,)


def self_update():
    """Check online for a newer vaani.py and, if there is one, swap it in
    and restart. Silent when offline or already up to date - the run just
    continues with the version on disk. Your leads are never uploaded;
    this only DOWNLOADS the program file."""
    if os.environ.get("VAANI_NO_UPDATE") == "1" or "--no-update" in sys.argv:
        return
    try:
        req = urllib.request.Request(UPDATE_URL,
                                     headers={"User-Agent": BROWSER_UA})
        with urllib.request.urlopen(req, timeout=5) as resp:
            content = resp.read().decode("utf-8", "replace")
    except Exception:
        return  # no internet / blocked - run as-is
    m = re.search(r'^VERSION\s*=\s*"([0-9.]+)"', content, re.M)
    if not m or _version_tuple(m.group(1)) <= _version_tuple(VERSION):
        return
    if "Vaani Lead Engine" not in content or len(content) < 10000:
        return  # not a real vaani.py - don't touch anything
    try:
        compile(content, "vaani.py", "exec")  # must be valid Python
    except Exception:
        return
    try:
        target = os.path.abspath(__file__)
        new_file = target + ".new"
        bak_file = target + ".bak"
        with open(new_file, "w", encoding="utf-8", newline="\n") as f:
            f.write(content)
        with open(target, "r", encoding="utf-8") as f:
            old = f.read()
        with open(bak_file, "w", encoding="utf-8", newline="\n") as f:
            f.write(old)
        os.replace(new_file, target)
    except Exception:
        return
    print(f"\n  Updated Vaani to v{m.group(1)} - restarting with the "
          "new version...", flush=True)
    env = dict(os.environ, VAANI_NO_UPDATE="1")
    try:
        rc = subprocess.run([sys.executable, os.path.abspath(__file__)]
                            + sys.argv[1:], env=env).returncode
    except Exception:
        rc = 0
    sys.exit(rc)

BROWSER_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
              "AppleWebKit/537.36 (KHTML, like Gecko) "
              "Chrome/130.0.0.0 Safari/537.36")

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
JUNK_DOMAINS = ("example.", "sentry.", "wix.", "godaddy", "squarespace",
                "wordpress", "schema.org", "w3.org", "yourdomain",
                "email.com", "test.", "sample.")


# ---------------------------------------------------------------- utilities

# Indian STD (landline area) codes, without the trunk 0 - metros first,
# then the major cities. Used to tell a landline (0 + STD + number) apart
# from a mobile written with a trunk 0 (0 + 10-digit mobile).
STD_CODES = {
    "11", "22", "33", "44", "20", "79", "80", "40",
    "120", "121", "129", "135", "141", "161", "171", "172", "177", "181",
    "183", "191", "194", "231", "233", "240", "251", "253", "256", "257",
    "261", "265", "271", "281", "291", "322", "341", "343", "353", "361",
    "413", "416", "422", "427", "431", "452", "462", "471", "484", "487",
    "495", "512", "522", "532", "542", "551", "562", "571", "581", "591",
    "612", "621", "631", "641", "651", "654", "657", "671", "674", "712",
    "721", "724", "731", "734", "751", "755", "761", "771", "775", "788",
    "821", "824", "831", "832", "836", "861", "863", "866", "870", "877",
    "891",
}


def _is_std_landline(digits11):
    """True for 0 + known STD code + a landline-style subscriber number
    (subscriber numbers start 2-5; mobile numbers start 6-9)."""
    body = digits11[1:]
    for ln in (4, 3, 2):
        if body[:ln] in STD_CODES:
            sub = body[ln:]
            if sub and sub[0] in "2345":
                return True
    return False


def clean_phone(raw):
    """Return the 10-digit Indian mobile number, or None.

    Accepted: 9825012345, +91 98250 12345, and 098250 12345 (trunk 0 -
    the format most businesses use on Maps). Landlines (079-..., 011-...
    i.e. 0 + STD code + subscriber number) are dropped on purpose -
    they are useless for outreach.
    """
    if not raw:
        return None
    digits = re.sub(r"[^\d]", "", str(raw))
    if digits.startswith("91") and len(digits) == 12:   # +91 98250 12345
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        # 0 + 10 digits: either a trunk-0 mobile or an STD landline
        if digits[1] in "6789" and not _is_std_landline(digits):
            digits = digits[1:]
    if len(digits) == 10 and digits[0] in "6789":       # real mobile
        return digits
    return None                                          # landline / invalid


def split_phones(raw):
    """One Google Maps phone field can hold several numbers - split & clean."""
    if not raw:
        return []
    out = []
    for part in re.split(r"[,;/|]", str(raw)):
        mob = clean_phone(part)
        if mob and mob not in out:
            out.append(mob)
    return out


def split_emails(raw):
    """Clean up found emails: valid format only, junk dropped, deduped."""
    if not raw:
        return []
    out = []
    for e in EMAIL_RE.findall(str(raw).lower()):
        domain = e.rpartition("@")[2]
        if any(j in domain for j in JUNK_DOMAINS):
            continue
        if e not in out:
            out.append(e)
    return out[:3]


def slug(text):
    s = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return s or "leads"


# ---------------------------------------------------------------- scraper

class Blocked(Exception):
    """Google is rate-limiting this connection - stop politely."""


def import_playwright():
    try:
        from playwright.sync_api import sync_playwright
        return sync_playwright
    except ImportError:
        sys.exit(
            "\n  The browser tool (Playwright) is not installed yet.\n"
            "  Run this ONE command in this folder, then run vaani.py again:\n"
            "\n      py -m pip install -r requirements.txt\n")


def launch_browser(p, headless):
    """Use the user's own Chrome if present (no download); else the
    built-in browser; else download it once and retry."""
    args = ["--disable-blink-features=AutomationControlled"]
    kw_base = {"headless": headless, "args": args}
    proxy = (os.environ.get("HTTPS_PROXY") or os.environ.get("https_proxy")
             or os.environ.get("HTTP_PROXY") or os.environ.get("http_proxy"))
    if proxy:
        kw_base["proxy"] = {"server": proxy}
    for channel, label in (("chrome", "your Google Chrome"),
                           (None, "the built-in browser")):
        try:
            kw = dict(kw_base)
            if channel:
                kw["channel"] = channel
            browser = p.chromium.launch(**kw)
            print(f"  browser ready ({label})")
            return browser
        except Exception:
            continue
    print("  downloading the browser (one time only, a few minutes)...")
    subprocess.run([sys.executable, "-m", "playwright", "install",
                    "chromium"])
    try:
        return p.chromium.launch(**kw_base)
    except Exception:
        sys.exit(
            "\n  Could not start a browser. Run this once, then try again:\n"
            "\n      py -m playwright install chromium\n")


def goto_retry(page, url, tries=3):
    """Open a page, retrying on flaky network errors instead of dying."""
    last = None
    for attempt in range(tries):
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=60000)
            return
        except Exception as e:
            last = e
            if attempt < tries - 1:
                page.wait_for_timeout(3000)
    raise last


def handle_consent(page):
    """Dismiss the Google cookie popup if it appears."""
    for label in ("Reject all", "Accept all"):
        try:
            btn = page.get_by_role("button", name=label)
            if btn.count():
                btn.first.click(timeout=3000)
                page.wait_for_timeout(1500)
                return
        except Exception:
            pass


def is_blocked(page):
    try:
        if "/sorry/" in page.url:
            return True
        body = page.locator("body").inner_text(timeout=3000).lower()
        return "unusual traffic" in body or "automated queries" in body
    except Exception:
        return False


def collect_links(page, query, max_results):
    """Open the Maps search, scroll the results list, gather place links."""
    url = "https://www.google.com/maps/search/" + urllib.parse.quote(query)
    goto_retry(page, url)
    page.wait_for_timeout(2500)
    handle_consent(page)
    if is_blocked(page):
        raise Blocked
    try:
        page.wait_for_selector('div[role="feed"], h1', timeout=20000)
    except Exception:
        pass
    feed = page.locator('div[role="feed"]')
    if feed.count() == 0:
        # the search went straight to a single place page
        return [page.url] if page.locator("h1").count() else []

    links, seen, stale = [], set(), 0
    while len(links) < max_results and stale < 4:
        found = page.eval_on_selector_all(
            'div[role="feed"] a[href*="/maps/place/"]',
            "els => els.map(e => e.href)")
        new = 0
        for href in found:
            key = href.split("?")[0]
            if key not in seen:
                seen.add(key)
                links.append(href)
                new += 1
        stale = 0 if new else stale + 1
        if page.locator("text=You've reached the end of the list").count():
            break
        try:
            feed.evaluate("el => el.scrollTo(0, el.scrollHeight)")
        except Exception:
            break
        page.wait_for_timeout(random.randint(900, 1600))
    return links[:max_results]


def text_of(locator, timeout=5000):
    try:
        el = locator.first
        el.wait_for(state="attached", timeout=timeout)
        return el.inner_text(timeout=timeout).strip()
    except Exception:
        return ""


def clean_field(text):
    """Detail fields can start with an icon character - strip it."""
    return re.sub(r"^[^\w]+", "", text or "").strip()


PHONE_IN_TEXT = re.compile(
    r"\+?91[\s-]?\d{5}[\s-]?\d{5}|0\d{2,4}[\s-]?\d{6,8}|\d{5}[\s-]?\d{5}")


def best_mobile_in(text):
    """The mobile that appears most often in a blob of text/HTML - on a
    place page that is the place's own number (it repeats in the data,
    the call link and the button text). Only valid mobiles count."""
    counts, best, best_n = {}, "", 0
    for cand in PHONE_IN_TEXT.findall(text or ""):
        mob = clean_phone(cand)
        if mob:
            counts[mob] = counts.get(mob, 0) + 1
            if counts[mob] > best_n:
                best, best_n = mob, counts[mob]
    return best


def _phone_candidates(page):
    """Yield phone-number candidates from the page, most reliable first.
    Verified on a live page: the row is a button named "Phone: ..." and
    contains a "Call phone number" link whose href is tel:..."""
    # 0) give the phone row a chance to appear at all (it can load late)
    try:
        page.wait_for_selector(
            'a[href^="tel:"], [data-item-id^="phone:tel:"], '
            'button[aria-label^="Phone:"]', timeout=15000)
    except Exception:
        pass
    # 1) the "Call phone number" link carries the number in its href
    try:
        a = page.locator('a[href^="tel:"]').first
        href = a.get_attribute("href", timeout=2000) or ""
        if href.startswith("tel:") and href[4:].strip():
            yield href[4:].strip()
    except Exception:
        pass
    # 2) the phone button, found by its accessible name "Phone: ..."
    try:
        b = page.get_by_role("button", name=re.compile(r"^Phone:")).first
        m = PHONE_IN_TEXT.search(b.inner_text(timeout=2000))
        if m:
            yield m.group(0)
    except Exception:
        pass
    # 3) the phone element's data attribute carries the number
    try:
        el = page.locator('[data-item-id^="phone:tel:"]').first
        item = el.get_attribute("data-item-id", timeout=2000) or ""
        if item.startswith("phone:tel:") and item[10:].strip():
            yield item[10:].strip()
        text = el.inner_text(timeout=1500).strip()
        if text:
            yield text
    except Exception:
        pass
    # 4) the button's aria-label attribute ("Phone: ...")
    try:
        el = page.locator('[aria-label^="Phone:"]').first
        label = el.get_attribute("aria-label", timeout=2000) or ""
        if ":" in label:
            yield label.split(":", 1)[1].strip()
    except Exception:
        pass
    # 5) scan the page's full HTML - the number is in the page data even
    #    when the panel itself renders oddly
    try:
        html = page.content()
        m = re.search(r'href="tel:([^"]+)"', html)
        if m:
            yield m.group(1)
        best = best_mobile_in(html)
        if best:
            yield best
    except Exception:
        pass
    # 6) last resort - all visible text on the page
    try:
        best = best_mobile_in(page.locator("body").inner_text(timeout=4000))
        if best:
            yield best
    except Exception:
        pass


def extract_phone(page):
    """First candidate that is a valid mobile wins. If none is a mobile,
    fall back to the first candidate found (e.g. a landline - the output
    filter drops it, so the place is excluded as intended)."""
    fallback = ""
    for cand in _phone_candidates(page):
        if not cand:
            continue
        if clean_phone(cand):
            return cand
        if not fallback:
            fallback = cand
    return fallback


def extract_rating_reviews(page):
    """Rating + review count from the place header. Verified structure:
    aria-label="4.8 stars" and aria-label="714 reviews" (the review
    breakdown rows use a different format, so these patterns are safe)."""
    rating, reviews = None, 0
    try:
        html = page.content()
    except Exception:
        return rating, reviews
    m = re.search(r'aria-label="(\d\.\d) stars"', html)
    if m:
        rating = float(m.group(1))
    m = re.search(r'aria-label="([\d,]+) reviews"', html)
    if m:
        reviews = int(m.group(1).replace(",", ""))
    return rating, reviews


def panel_text(page):
    try:
        return page.locator('div[role="main"]').first.inner_text(
            timeout=4000)
    except Exception:
        return ""


def extract_price_level(ptext):
    """Maps price band from the place header: ₹ = budget, ₹₹₹₹ = premium.
    Only the top of the side panel is scanned - the header (name,
    rating, category, price band) sits there, so a ₹ figure from a
    menu or review further down cannot be mistaken for the band.
    Returns the count of ₹ symbols, or None when Maps shows none."""
    head = ptext[:800]
    m = re.search(r"₹{1,4}", head)
    return len(m.group(0)) if m else None


def extract_place(page, url):
    """Read one restaurant's Maps page: name, address, phone, website."""
    goto_retry(page, url)
    try:
        page.wait_for_selector("h1", timeout=15000)
    except Exception:
        return None
    try:
        page.wait_for_selector(
            '[data-item-id="address"], [data-item-id^="phone:tel:"], '
            'a[data-item-id="authority"]', timeout=5000)
    except Exception:
        pass
    # scroll the side panel - lower rows (incl. phone) can load lazily
    try:
        page.mouse.move(250, 400)
        page.mouse.wheel(0, 1500)
        page.wait_for_timeout(350)
    except Exception:
        pass
    if is_blocked(page):
        raise Blocked
    website = ""
    try:
        w = page.locator('a[data-item-id="authority"]').first
        w.wait_for(state="attached", timeout=3000)
        website = w.get_attribute("href") or ""
    except Exception:
        pass
    rating, reviews = extract_rating_reviews(page)
    ptext = panel_text(page)
    low = ptext.lower()
    services = {k: (k in ptext)
                for k in ("Dine-in", "Takeaway", "Delivery")}
    price_level = extract_price_level(ptext)
    return {
        "name": text_of(page.locator("h1")),
        "address": clean_field(
            text_of(page.locator('button[data-item-id="address"]'))),
        "phone": extract_phone(page),
        "website": website,
        "rating": rating,
        "reviews": reviews,
        "category": text_of(page.locator('button[jsaction*="category"]')),
        "hours": clean_field(
            text_of(page.locator('[data-item-id="oh"]'))),
        "closed": ("permanently closed" in low
                   or "temporarily closed" in low),
        "services": services,
        "price_level": price_level,
    }


def fetch_site_info(website):
    """Everything useful from the business's own website, in one pass:
    emails, aggregator links (Zomato/Swiggy), direct-ordering capability,
    social links, WhatsApp, contact form, growth and call-heavy signals,
    and a proprietor/owner name ONLY if the site itself states one."""
    info = {"emails": [], "aggregators": [], "ordering": False,
            "socials": [], "whatsapp": False, "form": False,
            "growth": False, "call_heavy": False, "owner": ""}
    if not website or "google." in website:
        return info
    base = website.rstrip("/")
    pages = []
    for url in (website, base + "/contact", base + "/contact-us"):
        try:
            req = urllib.request.Request(
                url, headers={"User-Agent": BROWSER_UA})
            with urllib.request.urlopen(req, timeout=8) as resp:
                html = resp.read(300000).decode("utf-8", "replace")
        except Exception:
            continue
        pages.append(html)
        if not info["emails"]:
            info["emails"] = split_emails(html)
        if info["emails"] and len(pages) > 1:
            break
    if not pages:
        return info
    blob = "\n".join(pages)
    low = blob.lower()
    for h in re.findall(r'href="([^"]+)"', blob):
        hl = h.lower()
        if "zomato.com" in hl and "Zomato" not in info["aggregators"]:
            info["aggregators"].append("Zomato")
        if "swiggy.com" in hl and "Swiggy" not in info["aggregators"]:
            info["aggregators"].append("Swiggy")
        if "instagram.com" in hl and "Instagram" not in info["socials"]:
            info["socials"].append("Instagram")
        if "facebook.com" in hl and "Facebook" not in info["socials"]:
            info["socials"].append("Facebook")
        if "wa.me" in hl or "api.whatsapp.com" in hl:
            info["whatsapp"] = True
        if ("order" in hl and "zomato" not in hl and "swiggy" not in hl):
            info["ordering"] = True
    if "add to cart" in low or "checkout" in low:
        info["ordering"] = True
    info["form"] = "<form" in low
    plain = re.sub(r"<[^>]+>", " ", blob)
    info["growth"] = bool(re.search(
        r"franchise|our outlets|our locations|multiple locations",
        plain, re.I))
    info["call_heavy"] = bool(re.search(
        r"banquet|reserve a table|reservations|bulk orders|catering|"
        r"party orders", plain, re.I))
    m = re.search(r"(?:proprietor|owner)\s*[:\-\u2013]\s*"
                  r"([A-Z][a-z]+(?: [A-Z][a-z]+){0,2})", plain, re.I)
    if m:
        info["owner"] = m.group(1)
    return info


def fetch_emails(website):
    """Backwards-compatible wrapper: just the emails from a site."""
    return fetch_site_info(website)["emails"]


# ---------------------------------------------------------------- scoring
# Book My Food Phase 1 rubric (100 points, every point shows its evidence).
# The bands are starting estimates - they get recalibrated against real
# demo outcomes from the tracker, not treated as facts.

KNOWN_CHAINS = ("mcdonald", "domino", "pizza hut", "kfc", "burger king",
                "subway", "starbucks", "haldiram", "barbeque nation",
                "wow momo", "wow! momo", "la pino", "baskin", "cafe coffee",
                "jumbo king", "behrouz", "faasos", "freshmenu",
                "oven story", "chaayos", "third wave", "burger singh")


def score_lead(r):
    """Score one scraped place. Returns (score, tier, reasons, angle).
    Reasons are short evidence strings, best first - they land in the
    tracker so the caller sees WHY this lead ranks where it does."""
    web = r.get("web") or {}
    scored = []  # (points, reason)

    def add(pts, reason):
        if pts:
            scored.append((pts, reason))

    # Contactability - 20
    if split_phones(r.get("phone")):
        add(8, "Mobile number")
    if r.get("emails"):
        add(5, "Email found")
    if r.get("website"):
        add(4, "Has website")
    if web.get("socials"):
        add(3, "On " + "/".join(web["socials"]))
    contact = min(20, sum(x for x, _ in scored))
    # Activity - 20
    act = []
    rev = r.get("reviews") or 0
    if rev >= 500:
        act.append((8, f"{rev} reviews"))
    elif rev >= 200:
        act.append((6, f"{rev} reviews"))
    elif rev >= 50:
        act.append((4, f"{rev} reviews"))
    elif rev >= 10:
        act.append((2, f"{rev} reviews"))
    rating = r.get("rating")
    if rating:
        if 3.5 <= rating <= 4.7:
            act.append((6, f"Rated {rating}"))
        elif rating > 4.7:
            act.append((4, f"Rated {rating}"))
        elif rating >= 3.0:
            act.append((2, f"Rated {rating}"))
    if r.get("hours"):
        act.append((3, "Hours listed"))
    svcs = r.get("services") or {}
    if any(svcs.values()):
        act.append((3, "Services listed"))
    activity = min(20, sum(x for x, _ in act))
    scored += act
    # Product fit - 30
    fit = []
    if svcs.get("Delivery") and svcs.get("Takeaway"):
        fit.append((10, "Delivery + takeaway"))
    elif svcs.get("Delivery") or svcs.get("Takeaway"):
        fit.append((5, "Delivery or takeaway"))
    aggs = web.get("aggregators") or []
    if aggs:
        fit.append((10, "On " + "/".join(aggs)))
    if not web.get("ordering"):
        fit.append((10, "No direct ordering found"))
    fit_pts = min(30, sum(x for x, _ in fit))
    scored += fit
    # Pain - 20
    pain = []
    if aggs and not web.get("ordering"):
        pain.append((10, "Aggregator-dependent, no direct channel"))
    if not r.get("website"):
        pain.append((5, "No website"))
    elif not (r.get("emails") or web.get("form") or web.get("socials")):
        pain.append((2, "Thin website - no email, form or social"))
    if not web.get("ordering") and not aggs:
        pain.append((5, "No online ordering - manual/phone orders"))
    pain_pts = min(20, sum(x for x, _ in pain))
    scored += pain
    # Access / timing - 10
    acc = []
    name_low = (r.get("name") or "").lower()
    if not any(c in name_low for c in KNOWN_CHAINS):
        acc.append((5, "Independent / small group"))
    if web.get("growth"):
        acc.append((5, "Expansion / franchise signals"))
    acc_pts = min(10, sum(x for x, _ in acc))
    scored += acc

    total = contact + activity + fit_pts + pain_pts + acc_pts
    tier = "A" if total >= 75 else "B" if total >= 50 else "C"
    reasons = [t for _, t in sorted(scored, key=lambda x: -x[0])]
    if aggs and not web.get("ordering"):
        angle = "Aggregator-dependent"
    elif web.get("growth"):
        angle = "Growing business"
    elif not r.get("website"):
        angle = "Digitally weak - no website"
    elif web.get("call_heavy"):
        angle = "Call-heavy (reservations/banquets)"
    elif rating and rating >= 4.5 and rev >= 300:
        angle = "Reputation-active"
    elif not web.get("ordering"):
        angle = "No direct ordering"
    else:
        angle = "Operations upgrade"
    return total, tier, reasons, angle


# ---------------------------------------------------------------- delivery segment
# "Delivery-first mid-tier" (--segment delivery, restaurants only).
# The slice asked for 2026-10-24: places that live on delivery (and pay
# aggregator commission for it), have a real dining room - or are a
# cloud kitchen, which is delivery-only yet the most commission-exposed
# of all - are NOT famous or scaled, and are priced for everyday
# repeat orders (₹100-400 mains). Every rule is a public, checkable
# signal:
# "loyal base" is proxied by review count + rating, "commission pain"
# by aggregator presence / no direct ordering, and the price band by
# Maps' ₹ level (₹/₹₹ ~= budget-to-moderate; exact dish prices are
# not on Maps, so an unknown band is flagged, never guessed).
# The segment is an EXTRA file - it never removes anything from the
# normal outputs, the tracker or the master database.

SEGMENT_SMALL_CATS = ("cart", "stall", "kiosk", "hawker", "vendor")
SEGMENT_MIN_REVIEWS = 30
SEGMENT_MAX_REVIEWS = 1500
SEGMENT_MIN_RATING = 3.8
SEGMENT_FAMOUS_RATING = 4.7
SEGMENT_FAMOUS_REVIEWS = 500
SEGMENT_MAX_PRICE = 2  # ₹ or ₹₹ only; ₹₹₹/₹₹₹₹ = premium, excluded


def segment_check(lead):
    """Delivery-first mid-tier test. Returns (ok, reasons, bucket):
    reasons explain the verdict in plain words; bucket is the short
    exclusion label used for the summary counts."""
    svcs = lead.get("services") or {}
    name_low = (lead.get("name") or "").lower()
    cat_low = (lead.get("category") or "").lower()
    rev = lead.get("reviews") or 0
    rating = lead.get("rating")
    aggs = lead.get("web_aggs") or []
    ordering = lead.get("web_ordering")
    if not svcs.get("Delivery"):
        return False, ["No delivery service listed"], "no delivery"
    # Dining leg: a real dining room OR a cloud kitchen. Cloud kitchens
    # are delivery-only by design, yet they are the ones most exposed
    # to aggregator commission - and they survive on repeat orders
    # from a loyal audience, which is exactly this segment's target.
    # Detected from the Maps category or the business name.
    is_cloud = ("cloud kitchen" in name_low or "cloud kitchen" in cat_low
                or "ghost kitchen" in name_low
                or "ghost kitchen" in cat_low)
    if not svcs.get("Dine-in") and not is_cloud:
        return False, ["No dine-in and not a cloud kitchen"], "no dine-in"
    if any(c in name_low for c in KNOWN_CHAINS):
        return False, ["Chain / scaled brand"], "chain"
    if any(c in cat_low for c in SEGMENT_SMALL_CATS):
        return False, ["Cart / stall / kiosk - too small, no dining "
                       "or delivery setup"], "cart/small"
    if rev < SEGMENT_MIN_REVIEWS:
        return False, [f"Only {rev} reviews - no evidence of a loyal "
                       "base yet"], "too few reviews"
    if rating is not None and rating < SEGMENT_MIN_RATING:
        return False, [f"Rated {rating} - regulars are not coming "
                       "back"], "low rating"
    if rev > SEGMENT_MAX_REVIEWS:
        return False, [f"{rev} reviews - already too popular / "
                       "scaled"], "too popular"
    if (rating is not None and rating >= SEGMENT_FAMOUS_RATING
            and rev >= SEGMENT_FAMOUS_REVIEWS):
        return False, [f"Rated {rating} with {rev} reviews - famous, "
                       "dining-led brand"], "famous"
    if not aggs and ordering:
        return False, ["Has direct ordering and no aggregator link - "
                       "commission pain not visible"], "no commission pain"
    pl = lead.get("price_level")
    if pl is not None and pl > SEGMENT_MAX_PRICE:
        return False, ["₹" * pl + " on Maps - priced above the everyday "
                       "band"], "premium price"
    reasons = [("Cloud kitchen (delivery-only)" if is_cloud
                else "Delivery + dine-in"),
               f"{rev} reviews"]
    if rating is not None:
        reasons.append(f"Rated {rating}")
    if aggs:
        reasons.append("On " + "/".join(aggs))
    elif not ordering:
        reasons.append("No direct ordering found")
    if pl:
        reasons.append("₹" * pl + " on Maps")
    else:
        reasons.append("Price unknown - check menu")
    return True, reasons, ""


# ---------------------------------------------------------------- master DB
# One persistent record per business (SQLite, free, built into Python).
# Excel stays the calling interface; this is the memory underneath it -
# it stops the same restaurant being scraped and called as "new" forever,
# and it remembers who must never be contacted again.

MASTER_DB = "vaani_master.db"

# ------------------------------------------------------------------
# Verticals: restaurant data and healthcare data NEVER share a
# database. --vertical healthcare routes everything - the master
# database, the call tracker, the win-back list, the output folder -
# to healthcare-named files, so a clinic can never land inside the
# restaurant data (and the restaurant CRM never sees one).
# ------------------------------------------------------------------
VERTICALS = {
    "food": {
        "label": "Restaurants / Food",
        "db": "vaani_master.db",
        "tracker": "vaani-call-tracker.xlsx",
        "winback": "vaani-not-interested.xlsx",
        "name_header": "Restaurant name",
        "names_file": "restaurant_names.csv",
        "names_col": "restaurant_name",
        "count_label": "Restaurants in your database",
        "folder_prefix": "",
    },
    "healthcare": {
        "label": "Healthcare",
        "db": "vaani_healthcare.db",
        "tracker": "vaani-healthcare-tracker.xlsx",
        "winback": "vaani-healthcare-not-interested.xlsx",
        "name_header": "Clinic / hospital name",
        "names_file": "clinic_names.csv",
        "names_col": "clinic_name",
        "count_label": "Clinics / hospitals in your database",
        "folder_prefix": "healthcare_",
    },
}
VCFG = VERTICALS["food"]
WORKSPACE_KEY = "restaurants"
WORKSPACE_LABEL = "Restaurants"
HEALTHCARE_HINTS = ("clinic", "hospital", "dental", "dentist",
                    "doctor", "physio", "diagnostic", "pharmacy",
                    "nursing", "ayurved", "homeopath")


def apply_vertical(name):
    """Point every data file at the chosen vertical. Called once,
    from main(), before anything is read or written."""
    global MASTER_DB, TRACKER_NAME, WINBACK_NAME, VCFG
    global WORKSPACE_KEY, WORKSPACE_LABEL
    VCFG = VERTICALS[name]
    WORKSPACE_KEY = "healthcare" if name == "healthcare" else "restaurants"
    WORKSPACE_LABEL = "Healthcare" if name == "healthcare" else "Restaurants"
    MASTER_DB = VCFG["db"]
    TRACKER_NAME = VCFG["tracker"]
    WINBACK_NAME = VCFG["winback"]
    TRACKER_HEADERS[1] = VCFG["name_header"]


def apply_category(category, vertical="food"):
    """Route each scrape into its own CRM workspace automatically.

    Restaurants and healthcare retain their established databases. Any new
    category (Gyms, Schools, Salons, etc.) gets a separate SQLite database
    and tracker without requiring another CRM launcher.
    """
    text = str(category or "").strip()
    lowered = text.lower()
    if vertical == "healthcare" or any(hint in lowered for hint in HEALTHCARE_HINTS):
        apply_vertical("healthcare")
        return
    restaurant_words = (
        "restaurant", "restaurants", "food", "cafe", "cafes", "café",
        "delivery", "cloud kitchen", "eatery", "eateries", "dining",
    )
    if not text or any(word in lowered for word in restaurant_words):
        apply_vertical("food")
        return
    safe = re.sub(r"[^a-z0-9]+", "-", lowered).strip("-") or "other"
    label = " ".join(part.capitalize() for part in safe.split("-"))
    global MASTER_DB, TRACKER_NAME, WINBACK_NAME, VCFG
    global WORKSPACE_KEY, WORKSPACE_LABEL
    WORKSPACE_KEY = safe
    WORKSPACE_LABEL = label
    VCFG = {
        "label": label,
        "db": f"vaani_{safe}.db",
        "tracker": f"vaani-{safe}-tracker.xlsx",
        "winback": f"vaani-{safe}-not-interested.xlsx",
        "name_header": "Business name",
        "names_file": f"{safe}_names.csv",
        "names_col": "business_name",
        "count_label": f"{label} in your database",
        "folder_prefix": f"{safe}_",
    }
    MASTER_DB = VCFG["db"]
    TRACKER_NAME = VCFG["tracker"]
    WINBACK_NAME = VCFG["winback"]
    TRACKER_HEADERS[1] = VCFG["name_header"]

SUPPRESS_OUTCOMES = {"Not interested", "Wrong number", "Do not contact"}


def _master():
    db = sqlite3.connect(os.path.join(HERE, MASTER_DB))
    db.execute("""CREATE TABLE IF NOT EXISTS leads (
        phone TEXT PRIMARY KEY, name TEXT, address TEXT, website TEXT,
        email TEXT, rating REAL, reviews INTEGER, score INTEGER,
        tier TEXT, angle TEXT, first_seen TEXT, last_seen TEXT,
        times_seen INTEGER DEFAULT 1, suppressed INTEGER DEFAULT 0,
        suppress_reason TEXT)""")
    cols = [r[1] for r in db.execute("PRAGMA table_info(leads)")]
    if "suppressed_on" not in cols:
        db.execute("ALTER TABLE leads ADD COLUMN suppressed_on TEXT")
        db.commit()
    db.execute("CREATE TABLE IF NOT EXISTS crm_meta (key TEXT PRIMARY KEY, value TEXT)")
    db.execute(
        "INSERT OR REPLACE INTO crm_meta (key, value) VALUES (?, ?)",
        ("workspace_label", WORKSPACE_LABEL),
    )
    db.commit()
    return db


def _phone_keys(phone_field):
    keys = set()
    for part in re.split(r"[,;/|]", str(phone_field or "")):
        d = re.sub(r"\D", "", part)
        if len(d) >= 10:
            keys.add(d[-10:])
    return keys


def master_suppressed():
    try:
        db = _master()
        rows = db.execute(
            "SELECT phone FROM leads WHERE suppressed = 1").fetchall()
        db.close()
        return {r[0] for r in rows}
    except Exception:
        return set()


def master_suppress(phone_key, reason):
    try:
        db = _master()
        today = date.today().isoformat()
        db.execute(
            "INSERT INTO leads (phone, suppressed, suppress_reason,"
            " suppressed_on, first_seen, last_seen)"
            " VALUES (?, 1, ?, ?, ?, ?)"
            " ON CONFLICT(phone) DO UPDATE SET suppressed = 1,"
            " suppress_reason = excluded.suppress_reason,"
            " suppressed_on = COALESCE(leads.suppressed_on,"
            " excluded.suppressed_on)",
            (phone_key, reason, today, today, today))
        db.commit()
        db.close()
    except Exception:
        pass


def master_unsuppress(phone_key):
    try:
        db = _master()
        db.execute(
            "UPDATE leads SET suppressed = 0, suppress_reason = NULL,"
            " suppressed_on = NULL WHERE phone = ?", (phone_key,))
        db.commit()
        db.close()
    except Exception:
        pass


def master_upsert(lead):
    try:
        db = _master()
        today = date.today().isoformat()
        for key in _phone_keys(lead["phone"]):
            db.execute(
                "INSERT INTO leads (phone, name, address, website, email,"
                " rating, reviews, score, tier, angle, first_seen,"
                " last_seen, times_seen) VALUES (?,?,?,?,?,?,?,?,?,?,?,"
                "?,1) ON CONFLICT(phone) DO UPDATE SET name ="
                " excluded.name, address = excluded.address, website ="
                " excluded.website, email = excluded.email, rating ="
                " excluded.rating, reviews = excluded.reviews, score ="
                " excluded.score, tier = excluded.tier, angle ="
                " excluded.angle, last_seen = excluded.last_seen,"
                " times_seen = leads.times_seen + 1",
                (key, lead["name"], lead["location"], lead.get("website"),
                 lead["email"], lead.get("rating"), lead.get("reviews"),
                 lead.get("score"), lead.get("tier"), lead.get("angle"),
                 today, today))
        db.commit()
        db.close()
    except Exception:
        pass


def _crm_state():
    """Statuses set in the Vaani CRM app (its crm_state table), if
    the CRM has been used on this PC. For any phone the CRM knows,
    the CRM's word is final - the tracker no longer overrides it."""
    try:
        db = _master()
        has = db.execute(
            "SELECT 1 FROM sqlite_master WHERE type = 'table'"
            " AND name = 'crm_state'").fetchone()
        if not has:
            db.close()
            return {}
        rows = db.execute(
            "SELECT phone, status FROM crm_state").fetchall()
        db.close()
        return {r[0]: r[1] for r in rows}
    except Exception:
        return {}


def sync_suppressions(tracker_outcomes):
    """Merge the tracker's outcomes (and the CRM's, which win) into
    the master database, and return every phone that must stay out
    of this run's outputs."""
    crm_state = _crm_state()
    for k, info in tracker_outcomes.items():
        if k in crm_state:
            continue  # the CRM owns this lead now
        if info["outcome"] in SUPPRESS_OUTCOMES:
            master_suppress(k, info["outcome"])
        else:
            master_unsuppress(k)
    for k, status in crm_state.items():
        if status in SUPPRESS_OUTCOMES:
            master_suppress(k, status)
        else:
            master_unsuppress(k)
    out = master_suppressed()
    out |= {k for k, info in tracker_outcomes.items()
            if info["outcome"] in SUPPRESS_OUTCOMES
            and k not in crm_state}
    out |= {k for k, s in crm_state.items()
            if s in SUPPRESS_OUTCOMES}
    return out


def read_tracker_outcomes():
    """The user's latest word on every phone in the tracker:
    phone -> {outcome, name, location, email, notes}. Sheets are
    read in order, so a newer sheet's outcome overrides an older
    one - marking someone Interested again un-marks them."""
    path = os.path.join(HERE, TRACKER_NAME)
    if not os.path.exists(path) or not ensure_openpyxl():
        return {}
    import openpyxl
    out = {}
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            if ws.title == "How to use":
                continue
            headers = {}
            for c in range(1, (ws.max_column or 0) + 1):
                v = ws.cell(row=1, column=c).value
                if v in ("Phone", "Outcome", "Restaurant name",
                         "Location", "Email", "Notes") \
                        and v not in headers:
                    headers[v] = c  # first match wins (the summary
                    # block reuses the word "Outcome" further right)
            if "Phone" not in headers or "Outcome" not in headers:
                continue
            for r in range(2, (ws.max_row or 0) + 1):
                outcome = ws.cell(row=r,
                                  column=headers["Outcome"]).value
                if not outcome:
                    continue
                keys = _phone_keys(
                    ws.cell(row=r, column=headers["Phone"]).value)
                if not keys:
                    continue
                info = {"outcome": outcome}
                for field, col in (("name", "Restaurant name"),
                                   ("location", "Location"),
                                   ("email", "Email"),
                                   ("notes", "Notes")):
                    if col in headers:
                        info[field] = ws.cell(
                            row=r, column=headers[col]).value or ""
                for k in keys:
                    out[k] = dict(info)
        wb.close()
    except Exception:
        return {}
    return out


WINBACK_NAME = "vaani-not-interested.xlsx"


def _pretty_date(iso_text):
    try:
        return date.fromisoformat(str(iso_text)[:10]).strftime(
            "%d %b %Y")
    except Exception:
        return iso_text or ""


def export_winback(tracker_outcomes, crm_owned=frozenset()):
    """(Re)build vaani-not-interested.xlsx - every lead the user has
    marked Not interested, kept in one file so they can be contacted
    again later (a win-back list). Do Not Contact leads stay out
    forever and Wrong Number leads are dead numbers, so neither is
    ever included here."""
    merged = {}
    try:
        db = _master()
        rows = db.execute(
            "SELECT phone, name, address, email, score, tier, angle,"
            " suppressed_on, last_seen FROM leads WHERE suppressed = 1"
            " AND suppress_reason = 'Not interested'").fetchall()
        db.close()
        for r in rows:
            merged[r[0]] = {
                "name": r[1] or "", "location": r[2] or "",
                "email": r[3] or "", "score": r[4], "tier": r[5] or "",
                "angle": r[6] or "",
                "marked": _pretty_date(r[7] or r[8]), "notes": ""}
    except Exception:
        pass
    for key, info in tracker_outcomes.items():
        if key in crm_owned:
            continue  # the CRM's status for this lead is in the
            # database rows above - the tracker has no say anymore
        if info.get("outcome") != "Not interested":
            # the tracker has the final word - if the outcome changed
            # (Do Not Contact, or Interested again), the lead leaves
            # the win-back list too
            merged.pop(key, None)
            continue
        row = merged.setdefault(key, {
            "name": "", "location": "", "email": "", "score": None,
            "tier": "", "angle": "", "marked": "", "notes": ""})
        for f in ("name", "location", "email", "notes"):
            if info.get(f):
                row[f] = info[f]
    path = os.path.join(HERE, WINBACK_NAME)
    if not merged and not os.path.exists(path):
        return 0
    if not ensure_openpyxl():
        return 0
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Not interested"
    head_fill = PatternFill("solid", fgColor="1F4E5F")
    head_font = Font(bold=True, color="FFFFFF", size=12)
    for c, h in enumerate(["Restaurant Name", "Phone", "Location",
                           "Email ID", "Score", "Tier", "Pitch Angle",
                           "Marked Not Interested On", "Notes",
                           "Win-back Notes"], 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
    ordered = sorted(
        merged.items(),
        key=lambda kv: (-(kv[1]["score"] or 0), kv[1]["name"]))
    for r, (phone, row) in enumerate(ordered, 2):
        ws.cell(row=r, column=1, value=row["name"])
        ws.cell(row=r, column=2, value=phone)
        ws.cell(row=r, column=3, value=row["location"])
        ws.cell(row=r, column=4, value=row["email"])
        ws.cell(row=r, column=5, value=row["score"])
        ws.cell(row=r, column=6, value=row["tier"])
        ws.cell(row=r, column=7, value=row["angle"])
        ws.cell(row=r, column=8, value=row["marked"])
        ws.cell(row=r, column=9, value=row["notes"])
    widths = [38, 14, 55, 34, 7, 6, 24, 16, 50, 40]
    for i, w_ in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w_
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 26
    try:
        wb.save(path)
    except Exception:
        return 0
    return len(merged)


def scrape(location, category, max_results, want_email, headless):
    sync_playwright = import_playwright()
    query = f"{category} in {location}"
    results = []
    with sync_playwright() as p:
        browser = launch_browser(p, headless)
        ctx = browser.new_context(
            user_agent=BROWSER_UA, locale="en-IN",
            viewport={"width": 1366, "height": 800})
        page = ctx.new_page()
        try:
            def _block_heavy(route):
                # map tiles, photos and fonts are heavy and carry no lead
                # data - skipping them makes every page load much faster
                if route.request.resource_type in ("image", "media", "font"):
                    route.abort()
                else:
                    route.continue_()
            ctx.route("**/*", _block_heavy)
        except Exception:
            pass
        try:
            print(f'  searching Google Maps for "{query}" ...')
            links = collect_links(page, query, max_results)
            print(f"  found {len(links)} places - reading each one...\n")
            debug_saved = 0
            for i, url in enumerate(links, 1):
                try:
                    d = extract_place(page, url)
                except Blocked:
                    print("\n  Google is rate-limiting this connection - "
                          "stopping here.")
                    print("  (Wait 30-60 minutes before the next run.)")
                    break
                if not d or not d["name"]:
                    continue
                if not split_phones(d["phone"]) and debug_saved < 3:
                    # keep copies of missed pages - the user can send one
                    # for an exact fix instead of guesswork
                    try:
                        debug_saved += 1
                        name = ("debug_place_page.html" if debug_saved == 1
                                else f"debug_place_page_{debug_saved}.html")
                        with open(os.path.join(HERE, name),
                                  "w", encoding="utf-8") as f:
                            f.write(page.content())
                    except Exception:
                        pass
                d["emails"] = []  # websites are checked after, in parallel
                d["web"] = {}
                results.append(d)
                mob = split_phones(d["phone"])
                line = f"  [{i}/{len(links)}] {d['name'][:44]}"
                if d.get("rating"):
                    line += f"  - {d['rating']}* ({d.get('reviews', 0)} reviews)"
                line += f"  - mobile {mob[0]}" if mob else "  - no mobile"
                if d["emails"]:
                    line += f"  - {d['emails'][0]}"
                print(line)
                page.wait_for_timeout(random.randint(250, 600))
            if want_email:
                sites = [r for r in results if r.get("website")]
                if sites:
                    print(f"\n  checking {len(sites)} restaurant websites "
                          "for email IDs (several at a time)...")
                    with ThreadPoolExecutor(max_workers=6) as ex:
                        for r, inf in zip(sites, ex.map(
                                lambda x: fetch_site_info(x["website"]),
                                sites)):
                            r["web"] = inf
                            r["emails"] = inf["emails"]
                    got = sum(1 for r in sites if r["emails"])
                    print(f"  email IDs found for {got} of {len(sites)} "
                          "places")
            missed = sum(1 for r in results
                         if not split_phones(r["phone"]))
            if results and missed:
                print(f"\n  {missed} place(s) had no readable mobile "
                      "number.")
                print("  A copy of one such page is saved as "
                      "debug_place_page.html")
                print("  in this folder - send it to me if those places "
                      "should have")
                print("  numbers, and I will match the reader to the page.")
        except Blocked:
            print("\n  Google blocked the search (rate limit). "
                  "Wait 30-60 minutes and run again.")
        except KeyboardInterrupt:
            print("\n  Stopped by you - saving what we have so far.")
        except Exception:
            print("\n  The search hit a connection problem and stopped.")
            print("  Check your internet connection and run again.")
        finally:
            browser.close()
    return results


# ---------------------------------------------------------------- outputs

def ensure_openpyxl():
    try:
        import openpyxl  # noqa
        return True
    except ImportError:
        print("  installing the Excel helper (one time)...")
        subprocess.run([sys.executable, "-m", "pip", "install",
                        "openpyxl", "-q"])
        try:
            import openpyxl  # noqa
            return True
        except ImportError:
            print("  (!) Excel export skipped - couldn't install openpyxl. "
                  "CSVs are still saved.")
            return False


def write_csv(path, header, values):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([header])
        for v in values:
            w.writerow([v])


def write_excel(path, leads):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    head_fill = PatternFill("solid", fgColor="1F4E5F")
    head_font = Font(bold=True, color="FFFFFF", size=12)
    for c, h in enumerate(["Restaurant Name", "Location",
                           "Phone Number", "Email ID", "Score", "Tier",
                           "Pitch Angle", "Rating", "Reviews",
                           "Why this lead", "Date Assigned"], 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(vertical="center")
    for r, lead in enumerate(
            sorted(leads, key=lambda x: -(x.get("score") or 0)), 2):
        ws.cell(row=r, column=1, value=lead["name"])
        ws.cell(row=r, column=2, value=lead["location"])
        ws.cell(row=r, column=3, value=lead["phone"])
        ws.cell(row=r, column=4, value=lead["email"])
        ws.cell(row=r, column=5, value=lead.get("score"))
        ws.cell(row=r, column=6, value=lead.get("tier"))
        ws.cell(row=r, column=7, value=lead.get("angle"))
        ws.cell(row=r, column=8, value=lead.get("rating"))
        ws.cell(row=r, column=9, value=lead.get("reviews"))
        ws.cell(row=r, column=10, value=lead.get("why"))
        ws.cell(row=r, column=11, value=lead.get("assigned"))
    widths = [38, 55, 16, 34, 7, 6, 24, 8, 9, 60, 14]
    for i, w_ in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w_
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 26
    wb.save(path)


TRACKER_NAME = "vaani-call-tracker.xlsx"
OUTCOMES = ("Not called,No answer,Interested,Callback,"
            "Deck sent,Not interested,Wrong number")
FUNNEL_OUTCOMES = ("Not called,No answer,Staff reached,"
                   "Decision-maker reached,Conversation,Interested,"
                   "Demo booked,Demo done,Deck sent,Callback,"
                   "Not interested,Wrong number,Do not contact,Won,Lost")
SUMMARY_OUTCOMES = ["Not called", "No answer", "Interested", "Callback",
                    "Demo booked", "Demo done", "Won", "Not interested"]
TRACKER_HEADERS = ["#", "Restaurant name", "Phone", "Location", "Email",
                   "Score", "Tier", "Pitch angle", "Why this lead",
                   "Outcome", "Callback date", "Notes"]


def _style_tracker_sheet(ws, last_row):
    """The v3.6 tracker layout: lead columns A-L (with Score, Tier,
    Pitch angle and Why), the outcome summary at N/O, and the sales
    funnel as the Outcome dropdown in column J."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    head_fill = PatternFill("solid", fgColor="3B1F5E")
    head_font = Font(bold=True, color="FFFFFF")
    for c, h in enumerate(TRACKER_HEADERS, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
    for c, h in ((14, "Outcome"), (15, "Count")):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = head_fill
        cell.font = head_font
    for i, o in enumerate(SUMMARY_OUTCOMES):
        r = i + 2
        ws.cell(row=r, column=14, value=o)
        ws.cell(row=r, column=15, value=f"=COUNTIF(J2:J{last_row},N{r})")
    tr = len(SUMMARY_OUTCOMES) + 2
    ws.cell(row=tr, column=14, value="TOTAL")
    ws.cell(row=tr, column=15, value=f"=SUM(O2:O{tr - 1})")
    for col, w in {"A": 5, "B": 30, "C": 16, "D": 26, "E": 30, "F": 7,
                   "G": 6, "H": 24, "I": 52, "J": 17, "K": 13, "L": 30,
                   "N": 18, "O": 10}.items():
        ws.column_dimensions[col].width = w
    dv = DataValidation(type="list", formula1=f'"{FUNNEL_OUTCOMES}"',
                        allow_blank=True)
    dv.add(f"J2:J{last_row}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:L{last_row}"


def _read_sheet_rows(ws):
    """Read a dated sheet's leads in EITHER layout (v3.5: Outcome in F;
    v3.6: Score in F, Outcome in J) so old sheets upgrade in place
    without losing a single call outcome."""
    rows = []
    f1 = ws.cell(row=1, column=6).value
    for r in range(2, (ws.max_row or 0) + 1):
        name = ws.cell(row=r, column=2).value
        if not name:
            continue
        d = {"name": name, "phone": ws.cell(row=r, column=3).value,
             "location": ws.cell(row=r, column=4).value,
             "email": ws.cell(row=r, column=5).value, "score": None,
             "tier": None, "angle": None, "why": None, "outcome": None,
             "callback": None, "notes": None}
        if f1 == "Score":
            d.update(score=ws.cell(row=r, column=6).value,
                     tier=ws.cell(row=r, column=7).value,
                     angle=ws.cell(row=r, column=8).value,
                     why=ws.cell(row=r, column=9).value,
                     outcome=ws.cell(row=r, column=10).value,
                     callback=ws.cell(row=r, column=11).value,
                     notes=ws.cell(row=r, column=12).value)
        else:
            d.update(outcome=ws.cell(row=r, column=6).value,
                     callback=ws.cell(row=r, column=7).value,
                     notes=ws.cell(row=r, column=8).value)
        rows.append(d)
    return rows


def update_tracker(leads):
    """Put the leads straight into vaani-call-tracker.xlsx - no copy-paste.

    Leads go into a sheet named for today (e.g. "21 Oct 2026"), sorted
    best-first by score, each with its Tier, Pitch angle and Why. A run
    on another day creates that day's sheet in the SAME file; a second
    run the same day only adds leads whose phone is not already there.
    Sheets from older versions are upgraded in place, outcomes kept.
    Returns (added, already_there, sheet_name), or None if the tracker
    could not be updated (the CSV/Excel outputs are unaffected).
    """
    if not ensure_openpyxl():
        return None
    import openpyxl
    path = os.path.join(HERE, TRACKER_NAME)
    sheet_name = date.today().strftime("%d %b %Y")
    try:
        if os.path.exists(path):
            wb = openpyxl.load_workbook(path)
        else:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
    except Exception:
        return None
    try:
        preserved = []
        if sheet_name in wb.sheetnames:
            idx = wb.sheetnames.index(sheet_name)
            preserved = _read_sheet_rows(wb[sheet_name])
            wb.remove(wb[sheet_name])
            ws = wb.create_sheet(sheet_name, idx)
        else:
            ws = wb.create_sheet(sheet_name)
        existing = set()
        for d in preserved:
            existing |= _phone_keys(d.get("phone"))
        r = 1
        for d in preserved:
            r += 1
            _write_tracker_row(ws, r, d)
        added, already = 0, 0
        for lead in sorted(leads, key=lambda x: -(x.get("score") or 0)):
            keys = _phone_keys(lead["phone"])
            if keys & existing:
                already += 1
                continue
            r += 1
            _write_tracker_row(ws, r, lead)
            existing |= keys
            added += 1
        _style_tracker_sheet(ws, max(r, 41))
        wb.active = wb.sheetnames.index(sheet_name)
        wb.save(path)
        return added, already, sheet_name
    except PermissionError:
        print("\n  (!) The call tracker is open in Excel, so today's leads")
        print("      could not be added to it. Close vaani-call-tracker.xlsx")
        print("      and run again - the CSV and Excel files are saved.")
        return None
    except Exception as e:
        print(f"\n  (!) Could not update the call tracker ({e}).")
        print("      The CSV and Excel files are saved as usual.")
        return None


def _write_tracker_row(ws, r, d):
    ws.cell(row=r, column=1, value=r - 1)
    ws.cell(row=r, column=2, value=d.get("name"))
    ws.cell(row=r, column=3, value=d.get("phone"))
    ws.cell(row=r, column=4, value=d.get("location"))
    ws.cell(row=r, column=5, value=d.get("email"))
    ws.cell(row=r, column=6, value=d.get("score"))
    ws.cell(row=r, column=7, value=d.get("tier"))
    ws.cell(row=r, column=8, value=d.get("angle"))
    ws.cell(row=r, column=9, value=d.get("why"))
    ws.cell(row=r, column=10, value=d.get("outcome") or "Not called")
    ws.cell(row=r, column=11, value=d.get("callback"))
    ws.cell(row=r, column=12, value=d.get("notes"))


# ---------------------------------------------------------------- main

def main():
    # Update FIRST, parse arguments second. Otherwise a brand-new flag
    # (like --segment) is rejected by the OLD copy before the updater
    # ever gets a chance to fetch the version that knows it - the user
    # sees "unrecognized arguments" for a flag that already exists
    # upstream. self_update() needs no parsed arguments: it checks
    # sys.argv directly for --no-update and restarts with argv intact.
    self_update()
    ap = argparse.ArgumentParser(
        description=f"Vaani Lead Engine v{VERSION} - location in, 4 lead databases "
                    "out. Scraper built in.")
    ap.add_argument("--location", help='e.g. "Gota, Ahmedabad"')
    ap.add_argument("--category", default="restaurants",
                    help="what to search (default: restaurants)")
    ap.add_argument("--max", type=int, default=100,
                    help="max places to read (default: 100)")
    ap.add_argument("--no-email", dest="email", action="store_false",
                    help="skip website email lookup for a faster run")
    ap.add_argument("--headless", action="store_true",
                    help="hide the browser window while it works")
    ap.add_argument("--no-update", action="store_true",
                    help="skip the automatic update check")
    ap.add_argument("--vertical", choices=["food", "healthcare"],
                    default="food",
                    help="food (default) or healthcare - healthcare "
                         "leads go into their OWN database "
                         "(vaani_healthcare.db), never mixed with "
                         "restaurant data")
    ap.add_argument("--segment", choices=["delivery"], default=None,
                    help="delivery: also write delivery_leads.xlsx - "
                         "only delivery-first mid-tier places "
                         "(delivery + dine-in or cloud kitchen, not "
                         "famous, not carts, everyday prices)")
    a = ap.parse_args()
    requested_category = a.category
    apply_category(requested_category, a.vertical)
    if a.segment and WORKSPACE_KEY != "restaurants":
        print("  (Note: --segment only applies to restaurants, so it "
              "is ignored for this healthcare run.)")
        a.segment = None

    print(f"\n  VAANI LEAD ENGINE v{VERSION}")
    print("  ----------------------")
    print(f"  Vertical: {VCFG['label']}  (database: {MASTER_DB})")
    location = a.location or input(
        '\n  Type a location (e.g. "Gota, Ahmedabad"): ').strip()
    if not location:
        sys.exit("  No location given - nothing to search.")
    print()

    if a.vertical == "food" and WORKSPACE_KEY == "healthcare":
        print("  Category recognised as healthcare - using the Healthcare workspace.")
    elif WORKSPACE_KEY not in ("restaurants", "healthcare"):
        print(f"  New category recognised - using the separate {WORKSPACE_LABEL} workspace.")

    rows = scrape(location, a.category, a.max, a.email, a.headless)
    if not rows:
        sys.exit(
            "\n  No places were read. Either Google is rate-limiting right "
            "now\n  (wait 30-60 minutes and run again) or no listings "
            "matched.\n  Nothing was saved.\n")

    # ---- build the 4 databases -------------------------------------------
    # Rule: a restaurant with NO mobile number is not a lead - it is left
    # out of every file. (Landline-only counts as no number: landlines are
    # dropped by the phone filter, so nothing callable remains.)
    leads, seen_places = [], set()
    names, locations, phones, emails = [], [], [], []
    seen_phones, seen_emails = set(), set()
    skipped_no_phone = 0
    skipped_closed = 0
    for r in rows:
        key = (r["name"].lower(), r["address"].lower())
        if key in seen_places:
            continue
        seen_places.add(key)
        if r.get("closed"):
            skipped_closed += 1
            continue
        mob = split_phones(r["phone"])
        if not mob:
            skipped_no_phone += 1
            continue
        em = r.get("emails") or []
        names.append(r["name"])
        if r["address"]:
            locations.append(r["address"])
        for p in mob:
            if p not in seen_phones:
                seen_phones.add(p)
                phones.append(p)
        for e in em:
            if e not in seen_emails:
                seen_emails.add(e)
                emails.append(e)
        sc, tier, reasons, angle = score_lead(r)
        web = r.get("web") or {}
        leads.append({"name": r["name"], "location": r["address"],
                      "phone": ", ".join(mob), "email": ", ".join(em),
                      "website": r.get("website") or "",
                      "rating": r.get("rating"),
                      "reviews": r.get("reviews"),
                      "score": sc, "tier": tier, "angle": angle,
                      "why": "; ".join(reasons[:4]),
                      # Date Assigned: the run date, so a lead row
                      # always shows when it was added - new runs
                      # are instantly tellable from older leads.
                      "assigned": date.today().strftime("%d %b %Y"),
                      "services": r.get("services") or {},
                      "category": r.get("category") or "",
                      "price_level": r.get("price_level"),
                      "web_aggs": web.get("aggregators") or [],
                      "web_ordering": bool(web.get("ordering"))})

    # the tracker is the user's latest word: sync its outcomes into
    # the master database first (marking someone Interested again
    # un-marks them), then drop everything suppressed from the
    # outputs. Leads the Vaani CRM has a status for are owned by the
    # CRM - sync_suppressions lets the CRM win over the tracker.
    tracker_outcomes = read_tracker_outcomes()
    crm_owned = set(_crm_state())
    suppressed_phones = sync_suppressions(tracker_outcomes)
    kept, suppressed_n = [], 0
    for lead in leads:
        if _phone_keys(lead["phone"]) & suppressed_phones:
            suppressed_n += 1
            continue
        kept.append(lead)
    leads = kept

    if not leads:
        if suppressed_n:
            wb_n = export_winback(tracker_outcomes, crm_owned)
            msg = (
                "\n  Every lead from this search is already in your "
                "tracker marked\n  Not interested / Wrong number / Do "
                "not contact - nothing new to save.\n")
            if wb_n:
                msg += (
                    f"  Your Not interested leads are safe in "
                    f"{WINBACK_NAME} ({wb_n} lead(s)) - your win-back "
                    "list.\n")
            sys.exit(msg)
        sys.exit(
            f"\n  Found {len(rows)} places, but none of them had a mobile "
            "number\n  (landlines don't count). Nothing was saved - try a "
            "different area.\n")

    outdir = os.path.join(
        HERE, f"leads_{VCFG['folder_prefix']}{slug(location)}_"
              f"{date.today():%Y-%m-%d}")
    os.makedirs(outdir, exist_ok=True)
    write_csv(os.path.join(outdir, VCFG["names_file"]),
              VCFG["names_col"], names)
    write_csv(os.path.join(outdir, "locations.csv"), "location", locations)
    write_csv(os.path.join(outdir, "phone_numbers.csv"), "phone_number",
              phones)
    write_csv(os.path.join(outdir, "email_ids.csv"), "email_id", emails)

    excel_done = False
    if ensure_openpyxl():
        try:
            write_excel(os.path.join(outdir, "all_leads.xlsx"), leads)
            excel_done = True
        except Exception as e:
            print(f"  (!) Excel file failed: {e}")

    # ---- delivery-first mid-tier segment (extra file, nothing removed)
    seg_n = seg_price_known = 0
    seg_out = {}
    if a.segment == "delivery":
        seg_leads = []
        for lead in leads:
            ok, reasons, bucket = segment_check(lead)
            if ok:
                s = dict(lead)
                s["why"] = "; ".join(reasons)
                seg_leads.append(s)
            else:
                seg_out[bucket] = seg_out.get(bucket, 0) + 1
        seg_n = len(seg_leads)
        seg_price_known = sum(1 for l in seg_leads
                              if l.get("price_level"))
        if ensure_openpyxl():
            try:
                write_excel(os.path.join(outdir, "delivery_leads.xlsx"),
                            seg_leads)
            except Exception as e:
                print(f"  (!) Delivery segment file failed: {e}")

    for lead in leads:
        master_upsert(lead)
    tracker_result = update_tracker(leads)
    winback_n = export_winback(tracker_outcomes, crm_owned)

    # ---- summary -----------------------------------------------------------
    print("\n  DONE")
    print(f"  {VCFG['count_label']} : {len(names)} "
          "(every one has a mobile number)")
    print(f"  Left out - no mobile number  : {skipped_no_phone}")
    if skipped_closed:
        print(f"  Left out - closed business   : {skipped_closed}")
    if suppressed_n:
        print(f"  Left out - already marked    : {suppressed_n} "
              "(Not interested / Wrong number / Do not contact)")
    if winback_n:
        print(f"  Win-back list                : {winback_n} Not "
              f"interested lead(s) safe in {WINBACK_NAME}")
    if a.segment == "delivery":
        print(f"  Delivery segment             : {seg_n} lead(s) in "
              f"delivery_leads.xlsx ({seg_price_known} price-known, "
              f"{seg_n - seg_price_known} price-unknown - check menu)")
        for b, n in sorted(seg_out.items(), key=lambda x: -x[1]):
            print(f"    - left out ({b}): {n}")
    print(f"  Email IDs                    : {len(emails)}")
    tiers = {"A": 0, "B": 0, "C": 0}
    for l in leads:
        tiers[l.get("tier") or "C"] += 1
    print(f"  Lead quality                 : {tiers['A']} A - call today, "
          f"{tiers['B']} B, {tiers['C']} C")
    if tracker_result:
        t_added, t_already, t_sheet = tracker_result
        print(f"  Call tracker                 : {t_added} lead(s) added "
              f"to sheet '{t_sheet}'")
        if t_already:
            print(f"                               : {t_already} already in "
                  "today's sheet - skipped")
        print(f"                               : in {TRACKER_NAME} - ready "
              "to call, no copy-paste")
    print(f"\n  Your 4 databases are in:\n  {outdir}")
    print(f"    - {VCFG['names_file']}")
    print("    - locations.csv")
    print("    - phone_numbers.csv   (mobiles only)")
    print("    - email_ids.csv")
    if excel_done:
        print("    - all_leads.xlsx      (everything combined)")
    if a.vertical == "healthcare":
        print("\n  Saved under HEALTHCARE (vaani_healthcare.db) - "
              "your restaurant")
        print("  data was not touched.")
    print()


if __name__ == "__main__":
    main()
